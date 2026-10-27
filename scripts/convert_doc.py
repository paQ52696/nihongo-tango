#!/usr/bin/env python3
"""
期中考單字庫（Word .docx 或一般表格式 PDF）→ CSV，再到 App 的「匯入」頁用欄位對應匯入。

還沒拿到期中考檔案，所以這支腳本做成「通用」版本：
  - .docx：優先讀表格（每個表格的每一列）；沒有表格時，讀段落並用 Tab／兩個以上空白／「｜」分欄
  - .pdf ：用 pdfplumber 讀表格；沒有表格時，逐行用空白分欄
  - 第一列若看起來是標題（含「漢字」「讀音」「意思」「課」等字），就當作欄位名稱；否則自動命名「欄1、欄2…」
  - 「第3課」「Lesson 3」「L3」這種單獨一行的標題，會被當成之後每一列的「課次」
輸出 UTF-8（含 BOM，Excel 直接開也不會亂碼）的 CSV，並印出成功／失敗筆數。

用法：
  python scripts/convert_doc.py data/期中考.docx -o data/out/midterm.csv
  python scripts/convert_doc.py data/期中考.pdf  -o data/out/midterm.csv
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

HEADER_HINTS = ('漢字', '讀音', '假名', '意思', '中文', '詞性', '例句', '課', '頁', 'kanji', 'reading', 'meaning')
LESSON_RE = re.compile(r'^\s*(第\s*\d+\s*[課课章回]|lesson\s*\d+|L\s*\d+|unit\s*\d+)\s*$', re.I)
SPLIT_RE = re.compile(r'\t|\s{2,}|｜|\|')
KANA_RE = re.compile(r'[぀-ヿ]')


def split_line(line: str) -> list[str]:
    return [c.strip() for c in SPLIT_RE.split(line) if c.strip()]


def rows_from_docx(path: Path) -> list[list[str]]:
    try:
        import docx  # python-docx
    except ImportError:
        sys.exit('請先安裝 python-docx：pip install python-docx')
    d = docx.Document(str(path))
    rows: list[list[str]] = []
    if d.tables:
        for t in d.tables:
            for r in t.rows:
                cells = [c.text.strip() for c in r.cells]
                # 合併儲存格會重複出現，去掉相鄰重複
                dedup = [c for i, c in enumerate(cells) if i == 0 or c != cells[i - 1] or not c]
                rows.append(dedup)
    else:
        for p in d.paragraphs:
            if p.text.strip():
                rows.append(split_line(p.text) if not LESSON_RE.match(p.text) else [p.text.strip()])
    return rows


def rows_from_pdf(path: Path) -> list[list[str]]:
    try:
        import pdfplumber
    except ImportError:
        sys.exit('請先安裝 pdfplumber：pip install pdfplumber')
    rows: list[list[str]] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            if tables:
                for t in tables:
                    rows += [[(c or '').strip() for c in r] for r in t]
            else:
                for line in (page.extract_text() or '').splitlines():
                    rows.append(split_line(line) if not LESSON_RE.match(line) else [line.strip()])
    return rows


def convert(rows: list[list[str]]):
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return [], [], []
    header = None
    first = rows[0]
    if len(first) >= 2 and not KANA_RE.search(''.join(first)) and any(h in ' '.join(first).lower() for h in HEADER_HINTS):
        header = rows[0]
        rows = rows[1:]
    width = max([len(header or [])] + [len(r) for r in rows])
    header = (header or []) + [f'欄{i + 1}' for i in range(len(header or []), width)]
    out, failed = [], []
    lesson = ''
    for i, r in enumerate(rows, start=2):
        if len(r) == 1 and LESSON_RE.match(r[0]):
            lesson = re.sub(r'\s+', '', r[0])
            continue
        if not any(KANA_RE.search(c) for c in r):
            failed.append((i, r, '整列沒有假名，可能不是單字列'))
            continue
        out.append(r + [''] * (width - len(r)) + [lesson])
    return header + ['課次'], out, failed


def main(argv=None):
    ap = argparse.ArgumentParser(description='Word／PDF 單字表 → CSV')
    ap.add_argument('src', type=Path)
    ap.add_argument('-o', '--output', type=Path, default=None)
    a = ap.parse_args(argv)
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    ext = a.src.suffix.lower()
    if ext == '.docx':
        rows = rows_from_docx(a.src)
    elif ext == '.pdf':
        rows = rows_from_pdf(a.src)
    else:
        sys.exit(f'不支援 {ext}；CSV、Excel、JSON 請直接在 App 的匯入頁匯入（.doc 請先用 Word 另存成 .docx）')
    header, out, failed = convert(rows)
    dst = a.output or a.src.with_suffix('.csv')
    dst.parent.mkdir(parents=True, exist_ok=True)
    with dst.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(out)
    print(f'成功 {len(out)} 列、略過 {len(failed)} 列 → {dst}')
    for i, r, why in failed[:50]:
        print(f'  第 {i} 列略過（{why}）：{" | ".join(r)}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
