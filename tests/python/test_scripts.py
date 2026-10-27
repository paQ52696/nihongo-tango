"""Python 腳本的測試：python -m unittest discover tests/python"""
import csv
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts'))

import convert_doc  # noqa: E402
import import_pdf  # noqa: E402


class TestImportPdfHelpers(unittest.TestCase):
    def test_parse_kanji(self):
        self.assertEqual(import_pdf.parse_kanji('開く(自V)')[:2], ('開く', ['自動詞']))
        self.assertEqual(import_pdf.parse_kanji('（お爺さん）')[:2], ('お爺さん', ['少用漢字']))
        self.assertEqual(import_pdf.parse_kanji('色々（な）')[0], '色々')

    def test_parse_example_notes(self):
        ex, notes = import_pdf.parse_example('船で行くと、一日かかります。（與ついたち同字異音）')
        self.assertEqual(ex, '船で行くと、一日かかります。')
        self.assertEqual(notes, ['與ついたち同字異音'])
        ex, notes = import_pdf.parse_example('廊下を走（はし）らないでください。')
        self.assertIn('走（はし）', ex)
        self.assertEqual(notes, [])

    def test_guess_pos(self):
        g = lambda r, k: import_pdf.guess_pos(r, k, import_pdf.parse_kanji(k)[0])[0]
        self.assertEqual(g('あく', '開く(自V)'), '動詞')
        self.assertEqual(g('あおい', '青い'), 'い形容詞')
        self.assertEqual(g('いや（な）', '嫌（な）'), 'な形容詞')
        self.assertEqual(g('アパート', ''), '名詞')
        self.assertEqual(g('おてあらい', 'お手洗い'), '名詞')
        self.assertEqual(g('～えん', '～円'), '接尾詞')

    @unittest.skipUnless((ROOT / 'data' / 'n5.pdf').exists(), '沒有 data/n5.pdf')
    def test_real_pdf(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / 'n5.json'
            import_pdf.main([str(ROOT / 'data' / 'n5.pdf'), '-o', str(out)])
            words = json.loads(out.read_text(encoding='utf-8'))['words']
            self.assertEqual(len(words), 535)
            self.assertEqual(words[0]['reading'], 'あう')
            self.assertEqual(words[0]['page'], 1)


class TestConvertDoc(unittest.TestCase):
    def test_docx_with_lessons(self):
        import docx

        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / 'midterm.docx'
            doc = docx.Document()
            doc.add_paragraph('第1課')
            doc.add_paragraph('学校\tがっこう\t學校')
            doc.add_paragraph('先生\tせんせい\t老師')
            doc.add_paragraph('第2課')
            doc.add_paragraph('切手  きって  郵票')
            doc.add_paragraph('這一行不是單字')
            doc.save(src)
            out = Path(d) / 'midterm.csv'
            convert_doc.main([str(src), '-o', str(out)])
            rows = list(csv.reader(out.open(encoding='utf-8-sig')))
            self.assertEqual(rows[0][-1], '課次')
            self.assertEqual(rows[1], ['学校', 'がっこう', '學校', '第1課'])
            self.assertEqual(rows[3], ['切手', 'きって', '郵票', '第2課'])
            self.assertEqual(len(rows), 4)

    def test_docx_table_with_header(self):
        import docx

        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / 't.docx'
            doc = docx.Document()
            t = doc.add_table(rows=3, cols=3)
            for i, row in enumerate([['漢字', '讀音', '意思'], ['会う', 'あう', '見面'], ['青い', 'あおい', '藍色的']]):
                for j, v in enumerate(row):
                    t.cell(i, j).text = v
            doc.save(src)
            out = Path(d) / 't.csv'
            convert_doc.main([str(src), '-o', str(out)])
            rows = list(csv.reader(out.open(encoding='utf-8-sig')))
            self.assertEqual(rows[0], ['漢字', '讀音', '意思', '課次'])
            self.assertEqual(rows[2][:3], ['青い', 'あおい', '藍色的'])


if __name__ == '__main__':
    unittest.main()
